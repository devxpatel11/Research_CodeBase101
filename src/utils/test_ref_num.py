import os
import docx
import win32com.client
import pymupdf
from docx.oxml import parse_xml
from docx.oxml.ns import nsdecls

doc = docx.Document('deliverables/Access-Template-2024 (1).docx')
p = doc.add_paragraph()
pPr = parse_xml(f'<w:pPr {nsdecls("w")}><w:pStyle w:val="References"/><w:numPr><w:numId w:val="0"/></w:numPr><w:ind w:left="360" w:hanging="360"/></w:pPr>')
p._p.insert(0, pPr)
p.add_run('[1] Test Reference with manual bracket number.')

doc.save('deliverables/test_ref.docx')

word = win32com.client.Dispatch('Word.Application')
doc_w = word.Documents.Open(os.path.abspath('deliverables/test_ref.docx'))
doc_w.ExportAsFixedFormat(os.path.abspath('deliverables/test_ref.pdf'), 17)
doc_w.Close(False)
word.Quit()

pdf = pymupdf.open('deliverables/test_ref.pdf')
txt = pdf[-1].get_text()
for line in txt.split('\n'):
    if 'Test Reference' in line:
        print('RENDERED LINE:', repr(line))
