import docx
import sys

sys.stdout.reconfigure(encoding='utf-8')
doc = docx.Document('deliverables/update_v4_edited.docx')

print(f"Total paragraphs: {len(doc.paragraphs)}")
print(f"Total tables: {len(doc.tables)}")

for i, p in enumerate(doc.paragraphs):
    omml = len(p._element.xpath('.//m:oMath'))
    img = len(p._element.xpath('.//a:blip'))
    txt = p.text.strip().replace('\n', ' ')
    if txt or omml or img:
        print(f"P{i:03d} [omml={omml}, img={img}]: {txt[:90]}")
