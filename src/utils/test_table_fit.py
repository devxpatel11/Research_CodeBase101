import docx
from docx.shared import Inches, Pt
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.enum.table import WD_TABLE_ALIGNMENT
from docx.oxml import parse_xml
from docx.oxml.ns import nsdecls
import win32com.client
import pymupdf
import os

doc_s = docx.Document('deliverables/update_v4_edited.docx')
doc_t = docx.Document('deliverables/Access-Template-2024 (1).docx')

# Let's inspect the 3 data tables in doc_s
data_tables = [tbl for tbl in doc_s.tables if len(tbl.rows) > 1]
print(f'Found {len(data_tables)} data tables.')

for idx, tbl in enumerate(data_tables):
    print(f'Table {idx}: {len(tbl.rows)} rows, {len(tbl.columns)} cols')
    # Print header text
    hdr = [c.text.strip().replace('\n', ' ') for c in tbl.rows[0].cells]
    print('  Headers:', hdr)
