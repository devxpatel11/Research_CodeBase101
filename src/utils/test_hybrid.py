"""
Test Version B: Hybrid
- Proper Section 0 break (Section I starts in 2 cols on Page 1!)
- All 7 figures are strictly single-column (3.35 in, ZERO section breaks!)
- Tables I, II, III span across both columns (6.9 in) with continuous section breaks
"""

import docx
from docx.shared import Inches, Pt, RGBColor
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.enum.table import WD_TABLE_ALIGNMENT
from docx.oxml import parse_xml, OxmlElement
from docx.oxml.ns import nsdecls, qn
import win32com.client
import pymupdf
import os
import copy
import sys

sys.stdout.reconfigure(encoding='utf-8')

TEMPLATE_PATH = r"deliverables/Access-Template-2024 (1).docx"
SOURCE_PATH = r"deliverables/update_v4_edited.docx"
OUT_DOCX = r"deliverables/test_v3_hybrid.docx"
OUT_PDF = r"deliverables/test_v3_hybrid.pdf"

FIG_CONFIGS = {
    "Figure 1:": "figures/fig1_system_architecture.png",
    "Figure 3:": "figures/fig3_evidence_schema.png",
    "Figure 4:": "figures/fig4_temporal_protocol.png",
    "Figure 5:": "figures/fig5_data_profile.png",
    "Figure 6:": "figures/fig6_learning_curves.png",
    "Figure 7:": "figures/fig7_tradeoff_plot.png",
    "Figure 8:": "figures/fig8_ablation_results.png"
}

def set_p_style(p_elem, style_id):
    pPr = p_elem.find('{http://schemas.openxmlformats.org/wordprocessingml/2006/main}pPr')
    if pPr is None:
        pPr = parse_xml(f'<w:pPr {nsdecls("w")}/>')
        p_elem.insert(0, pPr)
    pStyle = pPr.find('{http://schemas.openxmlformats.org/wordprocessingml/2006/main}pStyle')
    if pStyle is None:
        pStyle = parse_xml(f'<w:pStyle {nsdecls("w")} w:val="{style_id}"/>')
        pPr.insert(0, pStyle)
    else:
        pStyle.set('{http://schemas.openxmlformats.org/wordprocessingml/2006/main}val', style_id)

def set_cell_margins(cell, top=30, bottom=30, left=40, right=40):
    tcPr = cell._tc.get_or_add_tcPr()
    tcMar = OxmlElement('w:tcMar')
    for m, val in [('top', top), ('bottom', bottom), ('left', left), ('right', right)]:
        node = OxmlElement(f'w:{m}')
        node.set(qn('w:w'), str(val))
        node.set(qn('w:type'), 'dxa')
        tcMar.append(node)
    tcPr.append(tcMar)

def suppress_numPr(p_elem):
    pPr = p_elem.find('{http://schemas.openxmlformats.org/wordprocessingml/2006/main}pPr')
    if pPr is None:
        pPr = parse_xml(f'<w:pPr {nsdecls("w")}/>')
        p_elem.insert(0, pPr)
    numPr = pPr.find('{http://schemas.openxmlformats.org/wordprocessingml/2006/main}numPr')
    if numPr is not None:
        pPr.remove(numPr)
    pPr.append(parse_xml(f'<w:numPr {nsdecls("w")}><w:numId w:val="0"/></w:numPr>'))

def create_continuous_1col_break():
    p = parse_xml(
        f'<w:p {nsdecls("w")}>'
        f'<w:pPr>'
        f'<w:sectPr>'
        f'<w:type w:val="continuous"/>'
        f'<w:cols w:space="720"/>'
        f'</w:sectPr>'
        f'</w:pPr>'
        f'</w:p>'
    )
    return p

def create_continuous_2col_break():
    p = parse_xml(
        f'<w:p {nsdecls("w")}>'
        f'<w:pPr>'
        f'<w:sectPr>'
        f'<w:type w:val="continuous"/>'
        f'<w:cols w:num="2" w:space="400"/>'
        f'</w:sectPr>'
        f'</w:pPr>'
        f'</w:p>'
    )
    return p

def format_ieee_table(table, col_widths, is_data_table=True, font_size=8.0):
    table.alignment = WD_TABLE_ALIGNMENT.CENTER
    tblPr = table._tbl.tblPr
    if is_data_table:
        borders = parse_xml(
            f'<w:tblBorders {nsdecls("w")}>'
            f'<w:top w:val="single" w:sz="6" w:space="0" w:color="000000"/>'
            f'<w:left w:val="none"/>'
            f'<w:bottom w:val="single" w:sz="6" w:space="0" w:color="000000"/>'
            f'<w:right w:val="none"/>'
            f'<w:insideH w:val="single" w:sz="4" w:space="0" w:color="D5D8DC"/>'
            f'<w:insideV w:val="none"/>'
            f'</w:tblBorders>'
        )
        tblPr.append(borders)
    else:
        borders = parse_xml(
            f'<w:tblBorders {nsdecls("w")}>'
            f'<w:top w:val="none"/><w:left w:val="none"/><w:bottom w:val="none"/><w:right w:val="none"/>'
            f'<w:insideH w:val="none"/><w:insideV w:val="none"/>'
            f'</w:tblBorders>'
        )
        tblPr.append(borders)
        
    for row_idx, row in enumerate(table.rows):
        for col_idx, cell in enumerate(row.cells):
            if col_idx < len(col_widths):
                cell.width = Inches(col_widths[col_idx])
            set_cell_margins(cell, top=30, bottom=30, left=40, right=40)
            if is_data_table and row_idx == 0:
                tcPr = cell._tc.get_or_add_tcPr()
                shd = parse_xml(f'<w:shd {nsdecls("w")} w:fill="EAECEE"/>')
                tcPr.append(shd)
                for p in cell.paragraphs:
                    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
                    for r in p.runs:
                        r.font.name = "Times New Roman"
                        r.font.size = Pt(font_size)
                        r.font.bold = True
            elif is_data_table:
                if row_idx % 2 == 1:
                    tcPr = cell._tc.get_or_add_tcPr()
                    shd = parse_xml(f'<w:shd {nsdecls("w")} w:fill="F8F9F9"/>')
                    tcPr.append(shd)
                for p in cell.paragraphs:
                    if col_idx == 0:
                        p.alignment = WD_ALIGN_PARAGRAPH.LEFT
                    else:
                        p.alignment = WD_ALIGN_PARAGRAPH.CENTER
                    for r in p.runs:
                        r.font.name = "Times New Roman"
                        r.font.size = Pt(font_size)

def build_hybrid():
    doc_t = docx.Document(TEMPLATE_PATH)
    doc_s = docx.Document(SOURCE_PATH)

    body_t = doc_t._body._element
    body_s = doc_s._body._element

    # 1. Section 0 (Title, Authors, Affiliation, Abstract, Index Terms)
    p_title = doc_t.paragraphs[2]
    p_title.text = ""
    p_title.add_run(doc_s.paragraphs[0].text.strip())
    p_title.style = "Paper Title"

    p_auth = doc_t.paragraphs[3]
    p_auth.text = ""
    p1_lines = doc_s.paragraphs[1].text.strip().split('\n')
    auth_lines = [l.strip().rstrip(',') for l in p1_lines[:-1] if l.strip()]
    affil_line = p1_lines[-1].strip()
    p_auth.add_run(", ".join(auth_lines))
    p_auth.style = "AU"

    p_aff = doc_t.paragraphs[4]
    p_aff.text = ""
    p_aff.add_run(affil_line)
    p_aff.style = "PI"

    # Remove extra affiliation paragraphs (originally 5, 6, 7)
    for _ in range(3):
        body_t.remove(doc_t.paragraphs[5]._element)

    # Now paragraph 6 is Abstract
    p_abs = doc_t.paragraphs[6]
    p_abs._element.clear_content()
    for child in doc_s.paragraphs[2]._element:
        tag = child.tag.split('}')[-1]
        if tag in ('r', 'oMath', 'oMathPara'):
            p_abs._element.append(copy.deepcopy(child))
    p_abs.style = "Abstract"

    # Paragraph 7 is Index Terms
    p_it = doc_t.paragraphs[7]
    p_it._element.clear_content()
    for child in doc_s.paragraphs[3]._element:
        tag = child.tag.split('}')[-1]
        if tag in ('r', 'oMath', 'oMathPara'):
            p_it._element.append(copy.deepcopy(child))
    p_it.style = "IT"

    # Paragraph 8 is sectPr for Section 0 (single col full width)
    # Paragraph 9 is final sectPr for Section 1 (2 cols continuous)
    # Remove everything after index 8 except final sectPr
    while len(body_t) > 10:
        body_t.remove(body_t[9])

    s_children = list(body_s)
    after_heading = False
    in_2col = True

    for s_idx, elem in enumerate(s_children):
        tag = elem.tag.split('}')[-1]
        if tag == 'sectPr':
            continue

        if tag == 'p':
            txt = elem.xpath('string(.)').strip()
            has_pic = bool(elem.xpath('.//a:blip'))
            has_math = bool(elem.xpath('.//m:oMath'))

            # Skip Section 0 items
            if s_idx in (0, 1, 2, 3):
                continue

            # Image paragraph (all images inside 2-col flow at 3.35 in width!)
            if has_pic:
                if not in_2col:
                    body_t.insert(len(body_t) - 1, create_continuous_2col_break())
                    in_2col = True

                next_txt = ""
                if s_idx + 1 < len(s_children):
                    next_txt = s_children[s_idx + 1].xpath('string(.)').strip()
                
                matched_key = None
                for fig_key in FIG_CONFIGS:
                    if next_txt.startswith(fig_key):
                        matched_key = fig_key
                        break
                
                if matched_key:
                    img_path = FIG_CONFIGS[matched_key]
                    if os.path.exists(img_path):
                        p_img = doc_t.add_paragraph()
                        p_img.alignment = WD_ALIGN_PARAGRAPH.CENTER
                        p_img.paragraph_format.space_before = Pt(6)
                        p_img.paragraph_format.space_after = Pt(2)
                        run_img = p_img.add_run()
                        run_img.add_picture(img_path, width=Inches(3.35))
                        body_t.insert(len(body_t) - 1, p_img._element)
                        after_heading = False
                    continue

            # Empty spacing line without math
            if not txt and not has_math:
                continue

            new_p_elem = copy.deepcopy(elem)

            # Major Headings:
            if any(txt.startswith(h) for h in [
                "I. INTRODUCTION", "II. RELATED WORK", "III. DATASETS AND EMPIRICAL AUDITS",
                "IV. PROPOSED METHODOLOGY AND MATHEMATICAL FORMULATION", "V. EXPERIMENTAL SETUP",
                "VI. RESULTS AND STATISTICAL ANALYSIS", "VII. ABLATION STUDIES AND SENSITIVITY ANALYSIS",
                "VIII. INDUSTRIAL APPLICATIONS & OPERATIONAL WORKLOAD", "IX. LIMITATIONS AND THREATS TO VALIDITY",
                "X. CONCLUSION AND FUTURE WORK", "DECLARATIONS AND COMPLIANCE STATEMENTS", "REFERENCES"
            ]):
                if not in_2col:
                    body_t.insert(len(body_t) - 1, create_continuous_2col_break())
                    in_2col = True
                set_p_style(new_p_elem, "H1")
                after_heading = True

            # Subsection Headings:
            elif any(txt.startswith(sub) for sub in [
                "A. Procurement Knowledge Graphs", "B. Supplier Discovery", "C. Procurement Document Validation",
                "D. Learning to Rank and Tabular", "A. Pre-Decision Evidence Masking", "B. Three-Outcome Mandatory",
                "C. LambdaMART Relevance", "D. Bounded Optional RCI", "A. Time-Respecting Chronological",
                "B. Candidate Generation", "C. Baseline Comparison", "A. Observed-Award Ranking",
                "B. Paired Differences", "C. False Clearance", "D. Cold-Start and Buyer",
                "A. Policy Weight Beta", "B. Missing Optional Evidence"
            ]):
                if not in_2col:
                    body_t.insert(len(body_t) - 1, create_continuous_2col_break())
                    in_2col = True
                set_p_style(new_p_elem, "H2_Cont")
                suppress_numPr(new_p_elem)
                after_heading = True

            # Table Title:
            elif txt.startswith("TABLE I:") or txt.startswith("TABLE II:") or txt.startswith("TABLE III:"):
                # Table I, II, III span across page (1 col)
                if in_2col:
                    body_t.insert(len(body_t) - 1, create_continuous_1col_break())
                    in_2col = False
                set_p_style(new_p_elem, "TableTitle")
                after_heading = False

            # Figure Caption:
            elif txt.startswith("Figure "):
                set_p_style(new_p_elem, "FigCaption")
                suppress_numPr(new_p_elem)
                after_heading = False

            # References:
            elif txt.startswith("[") and any(f"[{k}]" in txt[:5] for k in range(1, 41)):
                if not in_2col:
                    body_t.insert(len(body_t) - 1, create_continuous_2col_break())
                    in_2col = True
                set_p_style(new_p_elem, "References")
                suppress_numPr(new_p_elem)
                pPr = new_p_elem.find('{http://schemas.openxmlformats.org/wordprocessingml/2006/main}pPr')
                if pPr is not None:
                    ind = pPr.find('{http://schemas.openxmlformats.org/wordprocessingml/2006/main}ind')
                    if ind is not None:
                        pPr.remove(ind)
                    pPr.append(parse_xml(f'<w:ind {nsdecls("w")} w:left="360" w:hanging="360"/>'))
                after_heading = False

            # Regular Body Paragraph:
            else:
                if not in_2col:
                    body_t.insert(len(body_t) - 1, create_continuous_2col_break())
                    in_2col = True
                if after_heading:
                    set_p_style(new_p_elem, "PARA")
                    after_heading = False
                else:
                    set_p_style(new_p_elem, "PARAIndent")

            body_t.insert(len(body_t) - 1, new_p_elem)

        elif tag == 'tbl':
            tbl = docx.table.Table(elem, doc_s)
            rows = len(tbl.rows)
            cols = len(tbl.columns)

            new_tbl_elem = copy.deepcopy(elem)
            body_t.insert(len(body_t) - 1, new_tbl_elem)
            new_tbl = docx.table.Table(new_tbl_elem, doc_t)

            if cols == 2 and rows == 1:
                # Equation table: display OMML math (inside 2-col flow)
                # Cell 0: formula, Cell 1: number with noWrap!
                tcPr1 = new_tbl.rows[0].cells[1]._tc.get_or_add_tcPr()
                tcPr1.append(parse_xml(f'<w:noWrap {nsdecls("w")}/>'))
                format_ieee_table(new_tbl, col_widths=[2.80, 0.55], is_data_table=False)
                after_heading = False
            else:
                # Data Table: Table I, II, III spanning across page width 6.9 inches
                if cols == 6:
                    format_ieee_table(new_tbl, col_widths=[1.5, 1.2, 1.1, 1.1, 1.0, 1.0], is_data_table=True, font_size=8.0)
                else:
                    col_w = [6.9 / cols] * cols
                    format_ieee_table(new_tbl, col_widths=col_w, is_data_table=True, font_size=8.0)
                after_heading = False
                # Switch back to 2 columns after the wide table!
                if not in_2col:
                    body_t.insert(len(body_t) - 1, create_continuous_2col_break())
                    in_2col = True

    doc_t.save(OUT_DOCX)
    print(f"Saved {OUT_DOCX} ({os.path.getsize(OUT_DOCX):,} bytes)")

    # Export to PDF via Word COM
    try:
        word = win32com.client.Dispatch("Word.Application")
        doc_w = word.Documents.Open(os.path.abspath(OUT_DOCX))
        doc_w.ExportAsFixedFormat(os.path.abspath(OUT_PDF), 17)
        doc_w.Close(False)
        word.Quit()
        print(f"Exported {OUT_PDF} ({os.path.getsize(OUT_PDF):,} bytes)")
    except Exception as e:
        print(f"Word COM error: {e}")

if __name__ == "__main__":
    build_hybrid()
