import docx

doc_t = docx.Document('deliverables/Access-Template-2024 (1).docx')

for i, p in enumerate(doc_t.paragraphs):
    has_pic = bool(p._element.xpath('.//a:blip'))
    txt = p.text.strip()
    if has_pic or 'Fig' in p.style.name or 'Table' in p.style.name:
        print(f'P{i:3d} | style={p.style.name:15s} | has_pic={has_pic} | text={txt[:60]}')

for i, tbl in enumerate(doc_t.tables):
    print(f'Table {i}: rows={len(tbl.rows)}, cols={len(tbl.columns)}')
