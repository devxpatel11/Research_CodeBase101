import sys
import docx
import pymupdf

sys.stdout.reconfigure(encoding='utf-8')

doc_s = docx.Document('deliverables/update_v4_edited.docx')
doc_t = docx.Document('deliverables/update_v3_IEEE_Access.docx')

# 1. Text Invariance
s_text = [p.text.strip() for p in doc_s.paragraphs if p.text.strip()]
for tbl in doc_s.tables:
    for row in tbl.rows:
        for cell in row.cells:
            for p in cell.paragraphs:
                if p.text.strip():
                    s_text.append(p.text.strip())

t_text = [p.text.strip() for p in doc_t.paragraphs if p.text.strip()]
for tbl in doc_t.tables:
    for row in tbl.rows:
        for cell in row.cells:
            for p in cell.paragraphs:
                if p.text.strip():
                    t_text.append(p.text.strip())

missing = [item for item in s_text if not any(item in t or t in item for t in t_text)]
print('=== 1. INVARIANCE CHECK ===')
print(f'Total Source text items: {len(s_text)}')
print(f'Total Target text items: {len(t_text)}')
print(f'Missing items: {len(missing)}')
if not missing:
    print('>>> 100.0% TEXT INVARIANCE CONFIRMED! <<<')

# 2. PDF Page Analysis
pdf = pymupdf.open('deliverables/update_v3_IEEE_Access.pdf')
print(f'\n=== 2. PDF PAGE ANALYSIS (Total Pages: {len(pdf)}) ===')
for i, page in enumerate(pdf):
    blocks = page.get_text('blocks')
    text_blocks = [b for b in blocks if b[4].strip()]
    col1 = [b for b in text_blocks if b[0] < 280 and b[2] < 320]
    col2 = [b for b in text_blocks if b[0] > 280]
    span = [b for b in text_blocks if b[0] < 280 and b[2] > 350]
    print(f'Page {i+1:2d}: col1={len(col1):2d} | col2={len(col2):2d} | span={len(span):2d} | y_max={max([b[3] for b in text_blocks], default=0):.1f}')

# 3. Equation Numbers Check
print('\n=== 3. EQUATIONS CHECK ===')
all_pdf_text = ' '.join([p.get_text() for p in pdf])
for eq_num in ['(1)', '(2)', '(3)', '(4)', '(5)', '(6)', '(7)']:
    found = eq_num in all_pdf_text
    print(f'Equation {eq_num}: {"FOUND" if found else "NOT FOUND"}')

# 4. Figures & Captions Check
print('\n=== 4. FIGURES & CAPTIONS CHECK ===')
for fig in ['Figure 1:', 'Figure 3:', 'Figure 4:', 'Figure 5:', 'Figure 6:', 'Figure 7:', 'Figure 8:']:
    found = fig in all_pdf_text
    has_dup = ('FIGURE 1. Figure 1:' in all_pdf_text) or ('FIGURE 2. Figure 1:' in all_pdf_text)
    print(f'{fig}: FOUND={found}')

# 5. Leftover Template Text Check
print('\n=== 5. LEFTOVER TEMPLATE TEXT CHECK ===')
for bad in ['3Electrical Engineering Department', 'University of Colorado', 'First A. Author', 'Boulder, CO']:
    print(f'Checking for "{bad}": {"FOUND ERROR" if bad in all_pdf_text else "CLEAN"}')
