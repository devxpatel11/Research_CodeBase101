import docx

doc_t = docx.Document('deliverables/Access-Template-2024 (1).docx')
print("=== TEMPLATE TABLES AND IMAGES ===")
for i, t in enumerate(doc_t.tables):
    print(f"Table {i}: rows={len(t.rows)}, cols={len(t.columns)}")
    # check width
    col_widths = [cell.width.inches for cell in t.rows[0].cells if cell.width]
    print(f"   col widths: {col_widths}, total={sum(col_widths)}")

for i, p in enumerate(doc_t.paragraphs):
    if p._element.xpath('.//a:blip'):
        print(f"Template Image at P{i}")
