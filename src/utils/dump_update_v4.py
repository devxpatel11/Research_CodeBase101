import docx

doc = docx.Document('deliverables/update_v4_edited.docx')
print(f"Total paragraphs in update_v4_edited.docx: {len(doc.paragraphs)}")
with open('src/utils/dump_update_v4.txt', 'w', encoding='utf-8') as f:
    for i, p in enumerate(doc.paragraphs):
        f.write(f"--- P{i} ---\n{p.text}\n")
print("Dumped to src/utils/dump_update_v4.txt")
