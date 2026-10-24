import docx
import os

doc = docx.Document('deliverables/update_v4_edited.docx')

print("=== IMAGE PARAGRAPHS AND SURROUNDING TEXT IN update_v4_edited.docx ===")
for i, p in enumerate(doc.paragraphs):
    if p._element.xpath('.//a:blip'):
        prev_p = doc.paragraphs[i-1].text.strip() if i > 0 else "None"
        next_p = doc.paragraphs[i+1].text.strip() if i+1 < len(doc.paragraphs) else "None"
        print(f"Image at P{i}:")
        print(f"   Prev P: {prev_p[:60]}")
        print(f"   Next P: {next_p[:60]}")
