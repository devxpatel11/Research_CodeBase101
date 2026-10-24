import docx
import copy
from docx.oxml import parse_xml
from docx.oxml.ns import nsdecls

doc_t = docx.Document('deliverables/Access-Template-2024 (1).docx')
doc_s = docx.Document('deliverables/update_v4_edited.docx')

# Get paragraph 30 from source (which has inline OMML)
p30 = doc_s.paragraphs[30]
p30_xml = copy.deepcopy(p30._element)

# Ensure pPr exists
pPr = p30_xml.find('{http://schemas.openxmlformats.org/wordprocessingml/2006/main}pPr')
if pPr is None:
    pPr = parse_xml(f'<w:pPr {nsdecls("w")}/>')
    p30_xml.insert(0, pPr)

# Set pStyle to PARA
pStyle = pPr.find('{http://schemas.openxmlformats.org/wordprocessingml/2006/main}pStyle')
if pStyle is None:
    pStyle = parse_xml(f'<w:pStyle {nsdecls("w")} w:val="PARA"/>')
    pPr.insert(0, pStyle)
else:
    pStyle.set('{http://schemas.openxmlformats.org/wordprocessingml/2006/main}val', 'PARA')

# Append to doc_t
doc_t._body._element.append(p30_xml)
doc_t.save('deliverables/test_p_style.docx')
print("Successfully tested paragraph style modification!")
