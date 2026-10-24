import docx
import copy
import os

doc_t = docx.Document('deliverables/Access-Template-2024 (1).docx')
doc_s = docx.Document('deliverables/update_v4_edited.docx')

# Test: check if we can copy a table with OMML
tbl_eq1 = doc_s.tables[1] # Equation 1
tbl_xml = copy.deepcopy(tbl_eq1._element)

# Add to template body
doc_t._body._element.append(tbl_xml)
test_path = 'deliverables/test_eq_copy.docx'
doc_t.save(test_path)
print(f"Saved test to {test_path}, size: {os.path.getsize(test_path)} bytes")

# Verify that test_eq_copy has the OMML table
doc_v = docx.Document(test_path)
has_omml = bool(doc_v.tables[-1]._element.xpath('.//m:oMath'))
print(f"Verification: test document has OMML in last table: {has_omml}")
